from pathlib import Path

from rag.loaders.base import BaseLoader, Document
from rag.preprocessing import clean_text


class DocxLoader(BaseLoader):
    """Word documents. Paragraphs styled "Heading N" or "Title" become Markdown headings."""

    extensions = (".docx",)

    def load(self, path: Path) -> list[Document]:
        import docx  # python-docx; imported lazily so it is only needed for .docx sources

        lines = []
        for para in docx.Document(str(path)).paragraphs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            style = (para.style.name or "") if para.style else ""
            if style == "Title":
                lines.append(f"# {text}")
            elif style.startswith("Heading"):
                level = style.removeprefix("Heading").strip()
                depth = min(int(level), 3) + 1 if level.isdigit() else 2
                lines.append(f"{'#' * depth} {text}")
            else:
                lines.append(text)
            lines.append("")

        text = clean_text("\n".join(lines))
        return [Document(text=text, metadata={"source": path.name, "page": None})] if text else []
